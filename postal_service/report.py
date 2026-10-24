"""Формирование отчётов о расчётах: Word (.docx) и Excel (.xlsx)."""
from datetime import date, datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .analytics import summarize_history
from .items import ITEM_CLASSES
from .tariffs import DeliveryMethod, DeliverySpeed

REPORTS_DIR = Path(__file__).resolve().parent.parent / "reports"

HISTORY_HEADERS = [
    "№", "Тип", "Отправитель", "Получатель", "Расстояние, км", "Срочность",
    "Способ доставки", "Расчётная масса, кг", "Итого, руб.", "Срок, дн.",
    "Дата доставки",
]



def _prepare_dir(output_dir):
    directory = REPORTS_DIR if output_dir is None else Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _file_path(directory, extension):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    return directory / f"postal_report_{timestamp}.{extension}"


def _type_name(code):
    return ITEM_CLASSES[code].TYPE_NAME


def _item_details(item):
    """Особые параметры отправления в виде текста."""
    options = item.options()
    parts = []
    if options.get("with_notification"):
        parts.append("уведомление о вручении")
    if options.get("is_printed"):
        parts.append("печатная продукция")
    if options.get("declared_value", 0) > 0:
        parts.append(f"объявленная ценность {options['declared_value']:g} руб.")
    if options.get("fragile"):
        parts.append("хрупкое")
    return ", ".join(parts) if parts else "-"


def _history_row(row):
    delivery_date = date.fromisoformat(row["delivery_date"])
    return [
        row["id"],
        _type_name(row["item_type"]),
        row["sender"],
        row["recipient"],
        row["distance_km"],
        DeliverySpeed.from_code(row["speed"]).title,
        DeliveryMethod.from_code(row["delivery_method"]).title,
        row["billable_weight_kg"],
        row["total_cost"],
        row["delivery_days"],
        delivery_date.strftime("%d.%m.%Y"),
    ]


def export_word(items, history, output_dir=None):
    """Сохраняет отчёт в формате .docx и возвращает путь к файлу.

    items   - словарь {id: PostalItem};
    history - список словарей с расчётами (результат load_quotes()).
    """
    file_path = _file_path(_prepare_dir(output_dir), "docx")
    summary = summarize_history(history)

    document = Document()
    document.add_heading("Почтовые отправления: отчёт о расчётах", 0)
    document.add_paragraph(f"Отчёт сформирован: {datetime.now():%d.%m.%Y %H:%M}")

    document.add_heading("Общая статистика", level=1)
    document.add_paragraph(f"Количество отправлений: {len(items)}")
    document.add_paragraph(f"Количество расчётов: {summary['count']}")
    document.add_paragraph(f"Общая сумма расчётов: {summary['total_sum']:.2f} руб.")
    document.add_paragraph(f"Средняя стоимость: {summary['average']:.2f} руб.")
    document.add_paragraph(f"Минимальная стоимость: {summary['min']:.2f} руб.")
    document.add_paragraph(f"Максимальная стоимость: {summary['max']:.2f} руб.")

    document.add_heading("Отправления", level=1)
    if items:
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, header in zip(
            table.rows[0].cells,
            ["№", "Тип", "Отправитель", "Получатель", "Масса, кг", "Особенности"],
        ):
            cell.text = header
        for item in items.values():
            cells = table.add_row().cells
            cells[0].text = str(item.item_id)
            cells[1].text = item.type_name
            cells[2].text = item.sender
            cells[3].text = item.recipient
            cells[4].text = f"{item.weight_kg:g}"
            cells[5].text = _item_details(item)
    else:
        document.add_paragraph("Отправления отсутствуют.")

    document.add_heading("Выполненные расчёты", level=1)
    if history:
        table = document.add_table(rows=1, cols=len(HISTORY_HEADERS))
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, HISTORY_HEADERS):
            cell.text = header
        for row in history:
            cells = table.add_row().cells
            for cell, value in zip(cells, _history_row(row)):
                cell.text = f"{value:.2f}" if isinstance(value, float) else str(value)
    else:
        document.add_paragraph("Расчёты пока не выполнялись.")

    document.save(file_path)
    return file_path


def export_excel(items, history, output_dir=None):
    """Сохраняет отчёт в формате .xlsx и возвращает путь к файлу."""
    file_path = _file_path(_prepare_dir(output_dir), "xlsx")
    summary = summarize_history(history)

    workbook = Workbook()

    # Лист 1. Сводка
    sheet = workbook.active
    sheet.title = "Сводка"
    sheet.append(["Показатель", "Значение"])
    sheet.append(["Количество отправлений", len(items)])
    sheet.append(["Количество расчётов", summary["count"]])
    sheet.append(["Общая сумма, руб.", summary["total_sum"]])
    sheet.append(["Средняя стоимость, руб.", summary["average"]])
    sheet.append(["Минимальная стоимость, руб.", summary["min"]])
    sheet.append(["Максимальная стоимость, руб.", summary["max"]])
    sheet.append(["Дата отчёта", datetime.now().strftime("%d.%m.%Y %H:%M")])

    # Лист 2. Отправления
    sheet = workbook.create_sheet("Отправления")
    sheet.append(["№", "Тип", "Отправитель", "Получатель", "Масса, кг",
                  "Длина, см", "Ширина, см", "Высота, см", "Особенности"])
    for item in items.values():
        sheet.append([
            item.item_id, item.type_name, item.sender, item.recipient,
            item.weight_kg, item.length_cm, item.width_cm, item.height_cm,
            _item_details(item),
        ])

    # Лист 3. Расчёты
    sheet = workbook.create_sheet("Расчёты")
    sheet.append(HISTORY_HEADERS)
    for row in history:
        sheet.append(_history_row(row))

    # Оформление всех листов
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="28405A")

        for column in sheet.columns:
            width = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column
            )
            sheet.column_dimensions[column[0].column_letter].width = min(width + 3, 40)

    workbook.save(file_path)
    return file_path
