from docx import Document
from openpyxl import load_workbook

from postal_service.database import load_items, load_quotes, save_item, save_quote
from postal_service.items import Parcel
from postal_service.pricing import Quote
from postal_service.report import export_excel, export_word
from postal_service.tariffs import DeliveryMethod, DeliverySpeed


def prepare_database(tmp_path):
    db_path = tmp_path / "postal.db"
    parcel = Parcel("Иванов", "Орлов", 10, 50, 40, 30,
                    declared_value=5000, fragile=True)
    save_item(parcel, db_path)
    save_quote(Quote(parcel, 2000, DeliverySpeed.EXPRESS, DeliveryMethod.COURIER),
               db_path)
    return load_items(db_path), load_quotes(db_path)


def test_export_word(tmp_path):
    items, history = prepare_database(tmp_path)
    path = export_word(items, history, output_dir=tmp_path / "out")

    assert path.exists() and path.suffix == ".docx"

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Количество отправлений: 1" in text
    assert "Количество расчётов: 1" in text
    assert "Общая сумма расчётов: 2539.00 руб." in text

    assert len(document.tables) == 2
    assert document.tables[0].rows[1].cells[1].text == "Посылка"
    assert "2539.00" in [cell.text for cell in document.tables[1].rows[1].cells]


def test_export_word_empty(tmp_path):
    path = export_word({}, [], output_dir=tmp_path)
    assert path.exists()

    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    assert "Отправления отсутствуют." in text
    assert "Расчёты пока не выполнялись." in text


def test_export_excel(tmp_path):
    items, history = prepare_database(tmp_path)
    path = export_excel(items, history, output_dir=tmp_path / "out")

    assert path.exists() and path.suffix == ".xlsx"

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Сводка", "Отправления", "Расчёты"]

    summary = workbook["Сводка"]
    assert summary["B2"].value == 1
    assert summary["B3"].value == 1
    assert summary["B4"].value == 2539.0

    shipments = workbook["Отправления"]
    assert shipments["B2"].value == "Посылка"
    assert shipments["E2"].value == 10.0

    quotes = workbook["Расчёты"]
    assert quotes["B2"].value == "Посылка"
    assert quotes["F2"].value == "Экспресс"
    assert quotes["G2"].value == "Курьером до двери"
    assert quotes["I2"].value == 2539.0
    assert quotes["J2"].value == 2
    workbook.close()


def test_export_excel_empty(tmp_path):
    path = export_excel({}, [], output_dir=tmp_path)
    assert path.exists()

    workbook = load_workbook(path)
    assert workbook["Сводка"]["B2"].value == 0
    assert workbook["Сводка"]["B3"].value == 0
    workbook.close()
